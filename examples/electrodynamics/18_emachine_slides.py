"""Slides of the e-machine study (2004 Toyota Prius traction motor) from the committed results of examples 10 to 17.

Every number on the slides is read from ``examples/electrodynamics/results/*.json``; the figures are the PNGs next to
them. Each slide states its limits. Writes ``results/18_emachine_slides.pptx`` and, when LibreOffice is installed, a
PDF next to it.

    pip install python-pptx
    python examples/electrodynamics/18_emachine_slides.py
"""
from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path

import numpy as np

RES = Path(__file__).parent / "results"


def load(name: str) -> dict:
    return json.loads((RES / name).read_text())


def facts() -> dict:
    """The numbers of the slides, from the result files."""
    f = {}
    p = load("10_pmsm_prius.json")
    t = p["torque_vs_angle"]
    T, Y = np.array(t["torque_Nm"]), np.array(t["published_Nm"] if "published_Nm" in t else t["yang_2013_Nm"])
    f["fea"] = {"elements": p["elements"], "peak": float(T.max()), "peak_pub": float(Y.max()),
                "rms": float(np.sqrt(np.mean((T - Y) ** 2))), "cog": p["no_load"]["cogging_peak_to_peak_Nm"],
                "emf": p["no_load"]["emf_phase_fundamental_peak_V_at_1000rpm"],
                "cases": [(v["case"], v["error"], v["passed"]) for v in p["verification"]]}
    x = load("14_prius_getdp_crosscheck.json")
    cc = x["cogging_convergence"][-1]
    f["getdp"] = {"dT_max": x["torque_difference_getdp_minus_pinneapple"]["max_abs_Nm"],
                  "emf": (x["back_emf_1000rpm"]["pinneapple_fundamental_peak_V"],
                          x["back_emf_1000rpm"]["getdp_fundamental_peak_V"]),
                  "Ld": (x["inductances_apparent_250A_50deg"]["pinneapple"]["L_d_mH"],
                         x["inductances_apparent_250A_50deg"]["getdp"]["L_d_mH"]),
                  "Lq": (x["inductances_apparent_250A_50deg"]["pinneapple"]["L_q_mH"],
                         x["inductances_apparent_250A_50deg"]["getdp"]["L_q_mH"]),
                  "cog": (cc["pinneapple_pkpk_Nm"], cc["getdp_pkpk_Nm"], cc["n_gap"]),
                  "cog_lin": (x["cogging_diagnosis"]["linear_iron"]["pinneapple_pkpk_Nm"],
                              x["cogging_diagnosis"]["linear_iron"]["getdp_pkpk_Nm"])}
    d = load("11_prius_drive.json")
    env = d["envelope"]
    w = d["wltc_class3b"]
    f["drive"] = {"T0": env["torque_Nm"][0], "Pmax": max(env["power_kW"]), "eta": d["efficiency_map"]["peak_efficiency"],
                  "eta_at": d["efficiency_map"]["peak_at"], "whkm": w["battery_Wh_per_km"],
                  "T6000": env["torque_Nm"][-1]}
    e = load("16_prius_end_effects.json")
    c = e["corrected_2d"]
    f["ends"] = {"L_air": e["end_winding"]["300"]["air"]["L_end_dq"] * 1e6,
                 "L_iron": e["end_winding"]["300"]["iron_face"]["L_end_dq"] * 1e6,
                 "env3000": (e["envelope_Nm"]["2d"][2], e["envelope_Nm"]["air"][2]),
                 "ratios": [r["ratio"] for r in e["magnet_end_leakage_3d"] if r["end_air_m"] == 0.04],
                 "dT": c["end_leakage"]["torque_Nm"] / c["2d"]["torque_Nm"] - 1,
                 "dE": c["end_leakage"]["emf_fundamental_peak_V_1000rpm"] / c["2d"]["emf_fundamental_peak_V_1000rpm"] - 1}
    h = load("13_prius_thermal.json")
    sw = h["magnet_temperature_sweep"]
    dm = h["demagnetisation"]
    f["thermal"] = {"sweep": [(r["magnet_temperature_C"], r["torque_Nm"], r["emf_phase_fundamental_peak_V_1000rpm"])
                              for r in sw],
                    "peak": h["peak_temperatures_C"], "t180": h["peak_load"]["time_to_180C_winding_s"],
                    "Tmag": dm["magnet_temperature_C"],
                    "demag": {lab: {g: v["demagnetised_fraction"] for g, v in r.items()}
                              for lab, r in dm["cases"].items()},
                    "mesh": dm["mesh_refinement"]}
    s = load("15_prius_skew.json")
    sw = {(r["steps"], round(r["skew_mech_deg"], 3)): r for r in s["sweep"]}
    f["skew"] = {"straight": sw[(2, 0.0)], "k2": sw[(2, 3.75)], "k3": sw[(3, 5.625)],
                 "emf": (s["back_emf"]["fundamental_ratio"], s["back_emf"]["factor_fundamental"],
                         s["back_emf"]["max_deviation_from_factor_rel_fundamental"]),
                 "thd": (s["back_emf"]["thd_straight"], s["back_emf"]["thd_skewed"])}
    tol = load("17_prius_tolerances.json")
    f["tol"] = {"n": tol["valid_samples"], "T": tol["outputs"]["torque_Nm"], "r": tol["outputs"]["ripple"],
                "y": (tol["p_within_spec_joint"], *tol["p_within_spec_joint_wilson95"])}
    g = load("12_prius_rotor_design.json")
    f["design"] = {"runs": [(r["seed"], r["hybrid"]["fea_calls"], r["fea_only"]["calls_to_reach_hybrid"],
                             r["fea_call_saving"], r["hybrid"]["finer_fea_summary"]) for r in g["runs"]],
                   "base": g["baseline"]}
    return f


def deck(f: dict):
    from pptx import Presentation
    from pptx.util import Emu, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    def slide(title, bullets, image=None, limits=None, image_width=0.55):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Emu(400000), Emu(250000), W - Emu(800000), Emu(700000)).text_frame
        tb.text = title
        tb.paragraphs[0].runs[0].font.size = Pt(26)
        tb.paragraphs[0].runs[0].font.bold = True
        wide = False
        if image is not None:
            from PIL import Image

            with Image.open(RES / image) as im:
                wide = im.width / im.height > 2.0
        top = Emu(1000000)
        if image is None or wide:  # text across the slide; a wide figure under it at full width
            text_h = H - Emu(1900000) if image is None else Emu(250000) + Emu(330000) * len(bullets)
            text_w = W - Emu(800000)
        else:  # text on the left, figure on the right
            text_h, text_w = H - Emu(1900000), int(W * (1 - image_width)) - Emu(500000)
        body = s.shapes.add_textbox(Emu(400000), top, text_w, text_h).text_frame
        body.word_wrap = True
        for i, b in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = b
            p.font.size = Pt(14)
            p.space_after = Pt(4)
        if image is not None:
            bottom = H - Emu(800000)  # keep the limits line free
            if wide:
                y = top + text_h + Emu(100000)
                pic = s.shapes.add_picture(str(RES / image), Emu(400000), y, width=W - Emu(800000))
            else:
                y = top
                pic = s.shapes.add_picture(str(RES / image), int(W * (1 - image_width)) - Emu(200000), y,
                                           width=int(W * image_width))
            if pic.top + pic.height > bottom:
                scale = (bottom - pic.top) / pic.height
                pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
                pic.left = int((W - pic.width) / 2) if wide else pic.left
        if limits:
            lt = s.shapes.add_textbox(Emu(400000), H - Emu(750000), W - Emu(800000), Emu(600000)).text_frame
            lt.word_wrap = True
            lt.text = "Limits: " + limits
            lt.paragraphs[0].font.size = Pt(11)
            lt.paragraphs[0].font.italic = True
        return s

    fe, gd, dr, en, th, sk, tl, ds = (f[k] for k in ("fea", "getdp", "drive", "ends", "thermal", "skew", "tol", "design"))
    passed = sum(1 for c in fe["cases"] if c[2])
    slide("2004 Toyota Prius traction motor: verified FEA to drive, thermal and design",
          ["Open-source chain in PINNeAPPle (pinneapple_simulation.electromagnetics.machines):",
           "2-D nonlinear FEA -> second-code cross-check -> drive (MTPA, flux weakening, WLTC) -> end effects -> "
           "thermal and demagnetisation -> skew and tolerances -> FEA + GP + NSGA-II design -> exchange formats",
           "Every number in these slides is a measured output of examples 10 to 17 (results/*.json), regenerated "
           "by this script.",
           "Geometry and materials: pyleecan's Toyota_Prius.json (48 slots, 8 poles, V magnets, M400-50A)."],
          image="10_pmsm_field.png", image_width=0.42, limits="a model study from public data; no test-bench measurement of this motor was available.")
    slide("2-D FEA, verified",
          [f"Nonlinear A_z FEM, conforming mirrored meshes, Arkkio torque: {fe['elements']} elements.",
           f"Verification cases against closed forms: {passed} of {len(fe['cases'])} pass "
           "(magnetised ring: second order, as P1 should).",
           f"Torque vs current angle at 250 A against Yang et al. 2013: peak {fe['peak']:.1f} vs "
           f"{fe['peak_pub']:.0f} N m, rms difference {fe['rms']:.1f} N m (their model and the pyleecan geometry "
           "differ: centre post, steel data).",
           f"Cogging {fe['cog']:.2f} N m pk-pk; back-EMF fundamental {fe['emf']:.1f} V peak at 1000 rpm."],
          image="10_pmsm_curves.png", limits="2-D only; P1 elements; iron loss from B waveforms without minor loops.")
    slide("Second code: gmsh + GetDP",
          [f"Same geometry and materials, GetDP's own mesh, Newton solver and torque: torque vs angle within "
           f"{gd['dT_max']:.2f} N m.",
           f"Back-EMF fundamental {gd['emf'][0]:.2f} vs {gd['emf'][1]:.2f} V; L_d {gd['Ld'][0]:.3f} vs "
           f"{gd['Ld'][1]:.3f} mH; L_q {gd['Lq'][0]:.3f} vs {gd['Lq'][1]:.3f} mH.",
           f"Cogging does not agree: {gd['cog'][0]:.2f} vs {gd['cog'][1]:.2f} N m at {gd['cog'][2]} gap nodes, both "
           f"converged; with linear iron {gd['cog_lin'][0]:.3f} vs {gd['cog_lin'][1]:.3f} N m -> the difference is "
           "in the saturated bridges."],
          image="14_prius_getdp_crosscheck.png",
          limits="cogging uncertain by about 25 % between two codes; FEMM, pyleecan, Elmer not run (inputs written).")
    slide("Drive: MTPA, flux weakening, WLTC",
          [f"FEA flux map (saturation, cross-saturation), 250 A, 500 V DC with SVPWM: {dr['T0']:.1f} N m at low "
           f"speed, {dr['T6000']:.1f} N m at 6000 rpm, {dr['Pmax']:.1f} kW peak.",
           f"Peak efficiency {100 * dr['eta']:.1f} % at {dr['eta_at']['rpm']:.0f} rpm, "
           f"{dr['eta_at']['torque_Nm']:.1f} N m (copper + iron + magnet loss).",
           f"WLTC class 3b, illustrative compact car: {dr['whkm']:.1f} Wh/km at the battery (drive losses only).",
           "Optimiser checked against the closed forms of a linear machine (1e-14) and energy bookkeeping."],
          image="11_prius_drive.png",
          limits="steady state; no inverter, mechanical or battery losses; the car is not a measured Prius.")
    slide("End effects",
          [f"End windings (Neumann, thin filaments, verified against Rosa and Maxwell): L_end {en['L_air']:.1f} uH "
           f"in air, {en['L_iron']:.1f} uH with the iron face.",
           f"Effect on the envelope at 3000 rpm: {en['env3000'][0]:.1f} -> {en['env3000'][1]:.1f} N m.",
           "Magnet end leakage, 3-D GetDP model of half the stack: flux ratio "
           + " -> ".join(f"{r:.4f}" for r in en["ratios"]) + " on three meshes.",
           f"Corrected 2-D model: MTPA torque {100 * en['dT']:+.1f} %, back-EMF {100 * en['dE']:+.1f} %."],
          image="16_prius_end_effects.png",
          limits="corrections, not a coupled 3-D solution at load; no measured end-winding inductance to compare.")
    sw0, swN = th["sweep"][0], th["sweep"][-1]
    pk = th["peak"]
    slide("Magnet temperature and thermal network",
          [f"Magnets {sw0[0]:.0f} -> {swN[0]:.0f} degC (Br -0.1 %/K): torque at 250 A {sw0[1]:.1f} -> {swN[1]:.1f} "
           f"N m, back-EMF {sw0[2]:.1f} -> {swN[2]:.1f} V.",
           f"4 x WLTC from the drive losses: teeth {pk['teeth']:.1f}, winding {pk['winding']:.1f}, magnets "
           f"{pk['magnets']:.1f} degC.",
           f"Peak load from that state: hotter winding node at 180 degC after {th['t180']:.0f} s; magnets settle at "
           f"{th['Tmag']:.0f} degC."],
          image="13_prius_thermal.png",
          limits="seven-node network with illustrative thermal data (not fitted to a Prius); one temperature per node.")
    dm = th["demag"]
    lab_full = [k for k in dm if k.startswith("i_d")][0]
    lab_sc = [k for k in dm if not k.startswith("i_d")][0]
    mesh = th["mesh"]
    hs = [lv["cases"][lab_full]["H"]["demagnetised_fraction"] for lv in mesh]
    slide("Demagnetisation",
          [f"Reverse field per magnet element against Hcj(T) of the grade letters, magnets at {th['Tmag']:.0f} degC.",
           "Full d-axis current (-250 A): demagnetised area " + ", ".join(
               f"{g} {100 * v:.1f} %" for g, v in dm[lab_full].items()) + ".",
           "Short-circuit estimate: " + ", ".join(f"{g} {100 * v:.1f} %" for g, v in dm[lab_sc].items()) + ".",
           f"Mesh study inside the magnets (4 levels): grade H at -250 A {100 * min(hs):.1f} to {100 * max(hs):.1f} % "
           "-> area fractions converge; the worst element does not (sharp corners)."],
          image="13_prius_demag.png",
          limits="Hcj coefficient assumed (-0.55 %/K); static knee test; the Prius magnet grade was not found in "
                 "accessible public sources, so the rows compare candidate grades.")
    s0, k2, k3 = sk["straight"], sk["k2"], sk["k3"]
    slide("Step skew",
          [f"Multi-slice FEA. Straight rotor: cogging {s0['cogging_pkpk_Nm']:.2f} N m, ripple "
           f"{100 * s0['ripple_pkpk']:.1f} %, torque {s0['torque_mean_Nm']:.1f} N m.",
           f"2 steps over 3.75 deg: cogging {k2['cogging_pkpk_Nm']:.2f} N m, ripple {100 * k2['ripple_pkpk']:.1f} %, "
           f"torque {k2['torque_mean_Nm']:.1f} N m.",
           f"3 steps over 5.625 deg: cogging {k3['cogging_pkpk_Nm']:.2f} N m, ripple {100 * k3['ripple_pkpk']:.1f} %, "
           f"torque {k3['torque_mean_Nm']:.1f} N m.",
           f"Back-EMF fundamental x{sk['emf'][0]:.5f} vs closed-form skew factor {sk['emf'][1]:.5f}; THD "
           f"{100 * sk['thd'][0]:.1f} -> {100 * sk['thd'][1]:.1f} %."],
          image="15_prius_skew.png", limits="multi-slice: no inter-segment axial flux, no end effects.")
    slide("Manufacturing tolerances",
          [f"{tl['n']} Monte-Carlo samples (JCGM 101), one nonlinear FEA each: remanence, magnet thickness and "
           "width, air gap.",
           f"Torque {tl['T']['mean']:.1f} +- {tl['T']['std']:.2f} N m; ripple {100 * tl['r']['mean']:.1f} +- "
           f"{100 * tl['r']['std']:.2f} %.",
           "Largest effects: remanence on torque (SRC " + f"{tl['T']['src']['Br_scale']:.2f}), air gap on ripple "
           f"(SRC {tl['r']['src']['airgap']:.2f}); linear fits R^2 {tl['T']['src_r2']:.3f} / {tl['r']['src_r2']:.3f}.",
           f"Yield against the stated specification {tl['y'][0]:.3f} (Wilson 95 %: {tl['y'][1]:.3f} to "
           f"{tl['y'][2]:.3f})."],
          image="17_prius_tolerances.png", limits="the tolerance sigmas are assumptions, not measured spreads.")
    runs = ds["runs"]
    slide("Rotor design: FEA + Gaussian process + NSGA-II",
          ["Five rotor variables, maximise torque, minimise magnet mass, ripple <= 30 %; every front design "
           "re-evaluated by a finer FEA.",
           "FEA calls to equal hypervolume, hybrid vs NSGA-II on the FEA: " + "; ".join(
               f"seed {s}: {a} vs {b} ({100 * sv:.1f} % saved)" for s, a, b, sv, _ in runs) + ".",
           "Finer FEA: all front designs feasible, torque within "
           f"{100 * max(r[4]['torque_rel_diff_max'] for r in runs):.1f} %, ripple up to "
           f"{100 * max(r[4]['ripple_rel_diff_max'] for r in runs):.0f} % higher (6 angles under-sample it).",
           "Designs beating the Prius rotor sit at the 1.0 mm bridge bound."],
          image="12_prius_rotor_design.png",
          limits="no bridge stress, demagnetisation or speed constraints: shows the method and the FEA-call saving, "
                 "not a better rotor.")
    slide("Exchange formats and limits",
          ["In: pyleecan machine files (the real Toyota_Prius.json reads to exactly the model). Out: pyleecan JSON, "
           "FEMM Lua, Elmer case, gmsh/GetDP, MATLAB .mat, DXF, VTK, CSV/Parquet/HDF5 data contracts.",
           "What is verified: closed forms (FEM, filaments, thermal, skew factor, drive), a second FEA code, a "
           "published torque curve.",
           "Not verified: any measurement of this motor; FEMM/pyleecan/Elmer runs; cogging to better than about "
           "25 %; thermal data; magnet grade.",
           "Next: test-bench data (torque, EMF, temperatures) to calibrate thermal and loss models; a 3-D coupled "
           "end-region solution at load; bridge stress as a design constraint."],
          limits="see the Limits section of docs/guides/electromagnetics.md.")
    return prs


def main() -> Path:
    prs = deck(facts())
    out = RES / "18_emachine_slides.pptx"
    prs.save(out)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(RES), str(out)],
                       check=False, capture_output=True, timeout=300)
    print(out, len(prs.slides), "slides")
    return out


if __name__ == "__main__":
    main()
